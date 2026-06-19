#!/usr/bin/env python3
# -*- coding: utf-8 -*-
from __future__ import annotations
"""
integration/export_service.py — 多格式导出服务

支持: CSV、JSON、Excel (xlsx)、HTML 报表
符合数码印刷行业数据交换规范
"""
import csv
import json
import logging
from typing import List, Dict, Any, Optional
from pathlib import Path
from datetime import datetime

from utils.logger import get_logger

logger = get_logger(__name__)


class ExportService:
    """多格式导出服务"""

    @staticmethod
    def to_csv(data: List[Dict], output_path: str, headers: List[str] = None,
               encoding: str = 'utf-8-sig') -> str:
        """导出为 CSV"""
        if not data:
            logger.warning("导出数据为空")
            return ""
        h = headers or list(data[0].keys())
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        with open(output_path, 'w', newline='', encoding=encoding) as f:
            w = csv.DictWriter(f, fieldnames=h, extrasaction='ignore')
            w.writeheader()
            w.writerows(data)
        logger.info(f"CSV 导出完成: {output_path} ({len(data)} 条)")
        return output_path

    @staticmethod
    def to_json(data: Any, output_path: str, indent: int = 2) -> str:
        """导出为 JSON"""
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        with open(output_path, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=indent, default=str)
        logger.info(f"JSON 导出完成: {output_path}")
        return output_path

    @staticmethod
    def to_excel(data: List[Dict], output_path: str,
                 sheet_name: str = "数据", headers: List[str] = None) -> str:
        """导出为 Excel (xlsx)，需要 openpyxl"""
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        except ImportError:
            logger.warning("openpyxl 未安装，回退到 CSV 导出")
            csv_path = output_path.rsplit('.', 1)[0] + '.csv'
            return ExportService.to_csv(data, csv_path, headers)

        if not data:
            logger.warning("导出数据为空")
            return ""

        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name

        h = headers or list(data[0].keys())

        # 表头样式
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )

        # 写入表头
        for col_idx, key in enumerate(h, 1):
            cell = ws.cell(row=1, column=col_idx, value=key)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        # 写入数据
        for row_idx, row_data in enumerate(data, 2):
            for col_idx, key in enumerate(h, 1):
                value = row_data.get(key, "")
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")

        # 自动调整列宽
        for col_idx, key in enumerate(h, 1):
            max_len = len(str(key))
            for row in ws.iter_rows(min_row=2, max_row=min(ws.max_row, 100), min_col=col_idx, max_col=col_idx):
                for cell in row:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = min(max_len + 4, 50)

        # 冻结首行
        ws.freeze_panes = "A2"

        wb.save(output_path)
        logger.info(f"Excel 导出完成: {output_path} ({len(data)} 条)")
        return output_path

    @staticmethod
    def to_html_report(data: List[Dict], output_path: str,
                       title: str = "QHI 数据报表",
                       headers: List[str] = None) -> str:
        """导出为 HTML 报表"""
        if not data:
            logger.warning("导出数据为空")
            return ""

        h = headers or list(data[0].keys())
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        rows_html = ""
        for row_data in data:
            cells = "".join(f"<td>{row_data.get(k, '')}</td>" for k in h)
            rows_html += f"<tr>{cells}</tr>\n"

        header_html = "".join(f"<th>{k}</th>" for k in h)

        html = f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="UTF-8">
<title>{title}</title>
<style>
body {{ font-family: "Microsoft YaHei", sans-serif; margin: 20px; }}
h1 {{ color: #333; border-bottom: 2px solid #4CAF50; padding-bottom: 10px; }}
.meta {{ color: #666; font-size: 12px; margin-bottom: 20px; }}
table {{ border-collapse: collapse; width: 100%; }}
th {{ background: #4CAF50; color: white; padding: 10px; text-align: left; }}
td {{ border: 1px solid #ddd; padding: 8px; }}
tr:nth-child(even) {{ background: #f9f9f9; }}
tr:hover {{ background: #e8f5e9; }}
.summary {{ background: #f5f5f5; padding: 15px; border-radius: 5px; margin-bottom: 20px; }}
</style>
</head>
<body>
<h1>{title}</h1>
<div class="meta">生成时间: {now} | 数据量: {len(data)} 条</div>
<div class="summary">本报表由 QHI拼版处理器 自动生成</div>
<table>
<thead><tr>{header_html}</tr></thead>
<tbody>
{rows_html}
</tbody>
</table>
</body>
</html>"""

        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(html)
        logger.info(f"HTML 报表导出完成: {output_path} ({len(data)} 条)")
        return output_path

    @staticmethod
    def auto_export(data: List[Dict], output_path: str, **kwargs) -> str:
        """根据文件扩展名自动选择导出格式"""
        ext = Path(output_path).suffix.lower()
        if ext == '.csv':
            return ExportService.to_csv(data, output_path, **kwargs)
        elif ext == '.json':
            return ExportService.to_json(data, output_path, **kwargs)
        elif ext in ('.xlsx', '.xls'):
            return ExportService.to_excel(data, output_path, **kwargs)
        elif ext == '.html':
            return ExportService.to_html_report(data, output_path, **kwargs)
        else:
            logger.warning(f"不支持的导出格式: {ext}，回退到 CSV")
            return ExportService.to_csv(data, output_path, **kwargs)
